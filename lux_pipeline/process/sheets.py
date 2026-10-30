"""The curators' spreadsheet: same-as / different-from pairs and record fixes.

The spreadsheet has three tabs, each with a header row:

    Same As          A: URI, B: URI
    Different From   A: URI, B: URI
    Fixes            A-G: source, identifier, class, equivalent, path, operation, argument

`SheetImport` is the process, whatever the spreadsheet is kept in: `fetch`
reads the tabs and writes ``<data_dir>/sameAs/google.csv``,
``<data_dir>/differentFrom/google.csv`` and ``<data_dir>/xpath_fixes.json``
(read by every mapper); `load` adds the two CSVs to the merged source's
equivalent / different indexes. A subclass only says how to read a tab's
rows (`read_rows`):

    GoogleSheetImport   the Google Sheet named by <data_dir>/google_sheet_id.txt
    XlsxSheetImport     an Excel workbook (.xlsx)
    CsvSheetImport      a directory with one <tab name>.csv per tab

`sheet_import` picks one. Together, Google's fetch and load are
`google-sames-diffs.py` and the two `load-csv-map2.py` calls of
`run-all.sh`. The output file names keep ``google.csv`` whatever the
format: the indexes are loaded from those paths, and a file import
replaces the sheet's pairs rather than adding a second set.
"""

import csv
import json
import logging
import os

logger = logging.getLogger("lux_pipeline")

SAME_SHEET = "Same As"
DIFFERENT_SHEET = "Different From"
FIXES_SHEET = "Fixes"
FIX_FIELDS = ("source", "identifier", "class", "equivalent", "path", "operation", "argument")
FIRST_ROW = 2  # row 1 is the header


def csv_paths(data_dir):
    """``{"equivs": same-as CSV, "diffs": different-from CSV}``"""
    return {"equivs": os.path.join(data_dir, "sameAs", "google.csv"),
            "diffs": os.path.join(data_dir, "differentFrom", "google.csv")}


def width(last_col):
    """Columns from A to `last_col` (a single letter)."""
    return ord(last_col.upper()) - ord("A") + 1


def cell_text(value):
    """A cell as text: ``None`` is empty, a whole float is an integer."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def tidy_row(cells, last_col):
    """`cells` as text, cut to A..`last_col`, trailing empty cells dropped
    (as the Sheets API returns rows)."""
    row = [cell_text(c) for c in list(cells)[:width(last_col)]]
    while row and not row[-1].strip():
        row.pop()
    return row


class SheetImport:
    """Read the curators' spreadsheet and load it. Subclasses implement
    `read_rows`, and `describe` for the log; `open` and `close` bracket
    the reads."""

    def __init__(self, cfgs):
        self.configs = cfgs
        self.data_dir = cfgs.data_dir

    # --- what a format supplies --------------------------------------------

    def describe(self):
        return type(self).__name__

    def open(self):
        """Connect to or open the spreadsheet, before the first read."""
        return None

    def close(self):
        return None

    def read_rows(self, name, last_col):
        """The rows of tab `name` after its header, columns A to
        `last_col`, each a list of strings (trailing empty cells may be
        missing); ``None`` when the spreadsheet has no such tab."""
        raise NotImplementedError

    # --- the process -------------------------------------------------------

    def canonical_pairs(self, rows, name):
        """``[[a, b]]`` canonicalized; rows that are not two URIs, or that
        do not canonicalize, are logged and dropped."""
        canonicalize = self.configs.canonicalize
        out = []
        for n, row in enumerate(rows, FIRST_ROW):
            if len(row) < 2 or not row[0].strip() or not row[1].strip():
                if any(c.strip() for c in row):
                    logger.warning(f"{name} row {n}: not two URIs: {row}")
                continue
            a, b = row[0].strip(), row[1].strip()
            af, bf = canonicalize(a), canonicalize(b)
            if not af:
                logger.warning(f"{name} row {n}: failed to canonicalize {a}")
            elif not bf:
                logger.warning(f"{name} row {n}: failed to canonicalize {b}")
            else:
                out.append([af, bf])
        return out

    def parse_fixes(self, rows):
        """The Fixes tab as dicts of `FIX_FIELDS`; rows with fewer than six
        cells are logged and dropped, the seventh (argument) may be empty."""
        out = []
        for n, row in enumerate(rows, FIRST_ROW):
            if len(row) < 6:
                if any(c.strip() for c in row):
                    logger.warning(f"{FIXES_SHEET} row {n}: needs at least six cells: {row}")
                continue
            cells = [c.strip() for c in row[:7]] + [""] * (7 - len(row[:7]))
            out.append(dict(zip(FIX_FIELDS, cells)))
        return out

    def write_pairs(self, path, pairs):
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", newline="") as fh:
            csv.writer(fh).writerows(pairs)

    def fetch(self):
        """Read every tab, then write the two CSVs and ``xpath_fixes.json``.
        Returns ``{"equivs": n, "diffs": n, "fixes": n}``, ``None`` for a
        tab the spreadsheet does not have. An empty or missing tab leaves
        the previous file in place (an empty Fixes tab still writes an empty list, as the script did). Read
        errors propagate before anything is written, so a half-read
        spreadsheet never overwrites a file."""
        logger.info(f"reading {self.describe()}")
        self.open()
        try:
            tabs = {name: self.read_rows(name, last_col)
                    for name, last_col in ((DIFFERENT_SHEET, "B"), (SAME_SHEET, "B"),
                                           (FIXES_SHEET, "G"))}
        finally:
            self.close()

        paths = csv_paths(self.data_dir)
        counts = {}
        for which, name in (("diffs", DIFFERENT_SHEET), ("equivs", SAME_SHEET)):
            if tabs[name] is None:
                logger.warning(f"no {name} tab; {paths[which]} left as it was")
                counts[which] = None
                continue
            pairs = self.canonical_pairs(tabs[name], name)
            counts[which] = len(pairs)
            if pairs:
                self.write_pairs(paths[which], pairs)
                logger.info(f"{name}: {len(pairs)} pairs written to {paths[which]}")
            else:
                logger.warning(f"{name}: no rows; {paths[which]} left as it was")

        fixfn = os.path.join(self.data_dir, "xpath_fixes.json")
        if tabs[FIXES_SHEET] is None:
            logger.warning(f"no {FIXES_SHEET} tab; {fixfn} left as it was")
            counts["fixes"] = None
        else:
            fixes = self.parse_fixes(tabs[FIXES_SHEET])
            with open(fixfn, "w") as fh:
                fh.write(json.dumps(fixes))
            counts["fixes"] = len(fixes)
            logger.info(f"{FIXES_SHEET}: {len(fixes)} fixes written to {fixfn}")
        return counts

    def load(self):
        """Add the two CSVs to the merged source's indexes (adds, never
        clears: other CSVs are loaded into the same indexes). Returns the
        paths loaded."""
        loader = self.configs.results["merged"].get("indexLoader")
        if loader is None:
            raise RuntimeError("the merged source has no indexLoader to load the sheet into")
        loaded = []
        for which, path in csv_paths(self.data_dir).items():
            if not os.path.exists(path):
                logger.warning(f"{path} does not exist; nothing to load into {which}")
                continue
            logger.info(f"loading {path} into {which}")
            loader.load(path, which)
            loaded.append(path)
        return loaded


class GoogleSheetImport(SheetImport):
    """The Google Sheet whose id is in ``<data_dir>/google_sheet_id.txt``.
    Needs the optional Google client libraries (``pip install -e
    '.[google]'``), imported only when the sheet is opened. OAuth
    credentials are ``<data_dir>/credentials.json``; the token minted from
    them is kept in ``<data_dir>/token.json``."""

    SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
    PAGE = 500

    def __init__(self, cfgs, sheet=None):
        super().__init__(cfgs)
        # the Sheets API `spreadsheets()` resource; built by `open` unless given
        self.sheet = sheet
        self.sheet_id = None

    def describe(self):
        return f"Google Sheet {self.sheet_id or '(id in google_sheet_id.txt)'}"

    def open(self):
        with open(os.path.join(self.data_dir, "google_sheet_id.txt")) as fh:
            self.sheet_id = fh.read().strip()
        if self.sheet is None:
            self.sheet = self.client()

    def client(self):
        """The Sheets API ``spreadsheets()`` resource, authorised from the
        token, refreshing it or running the browser flow on
        ``credentials.json`` when it is missing or expired."""
        try:
            from google.auth.transport.requests import Request
            from google.oauth2.credentials import Credentials
            from google_auth_oauthlib.flow import InstalledAppFlow
            from googleapiclient.discovery import build
        except ImportError as e:
            raise RuntimeError(f"the Google client libraries are not installed ({e}); "
                               f"pip install -e '.[google]'") from e

        tokfn = os.path.join(self.data_dir, "token.json")
        credfn = os.path.join(self.data_dir, "credentials.json")
        creds = (Credentials.from_authorized_user_file(tokfn, self.SCOPES)
                 if os.path.exists(tokfn) else None)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                if not os.path.exists(credfn):
                    raise RuntimeError(f"no valid token in {tokfn} and no {credfn} "
                                       f"to authorise with")
                creds = InstalledAppFlow.from_client_secrets_file(
                    credfn, self.SCOPES).run_local_server(port=0)
            with open(tokfn, "w") as fh:
                fh.write(creds.to_json())
        return build("sheets", "v4", credentials=creds).spreadsheets()

    def read_rows(self, name, last_col):
        """A page of rows at a time until a page comes back empty."""
        rows = []
        start = FIRST_ROW
        while True:
            rng = f"{name}!A{start}:{last_col}{start + self.PAGE - 1}"
            values = (self.sheet.values().get(spreadsheetId=self.sheet_id, range=rng)
                      .execute().get("values", []))
            if not values:
                return rows
            rows.extend(values)
            start += self.PAGE


class XlsxSheetImport(SheetImport):
    """An Excel workbook with the three tabs. Needs openpyxl (``pip install
    -e '.[xlsx]'``), imported only when the workbook is opened."""

    def __init__(self, cfgs, path):
        super().__init__(cfgs)
        self.path = path
        self.workbook = None

    def describe(self):
        return f"workbook {self.path}"

    def open(self):
        try:
            import openpyxl
        except ImportError as e:
            raise RuntimeError(f"openpyxl is not installed ({e}); pip install -e '.[xlsx]'") from e
        self.workbook = openpyxl.load_workbook(self.path, read_only=True, data_only=True)

    def close(self):
        if self.workbook is not None:
            self.workbook.close()
            self.workbook = None

    def read_rows(self, name, last_col):
        if name not in self.workbook.sheetnames:
            return None
        ws = self.workbook[name]
        return [tidy_row(r, last_col)
                for r in ws.iter_rows(min_row=FIRST_ROW, max_col=width(last_col), values_only=True)]


class CsvSheetImport(SheetImport):
    """A directory with one CSV per tab, named after it (``Same As.csv``,
    ``Different From.csv``, ``Fixes.csv``), each with its header row: what
    a spreadsheet program's export of each tab gives."""

    def __init__(self, cfgs, path):
        super().__init__(cfgs)
        self.path = path

    def describe(self):
        return f"CSV directory {self.path}"

    def read_rows(self, name, last_col):
        fn = os.path.join(self.path, f"{name}.csv")
        if not os.path.exists(fn):
            return None
        with open(fn, newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
        return [tidy_row(r, last_col) for r in rows[FIRST_ROW - 1:]]


def sheet_import(cfgs, path=None):
    """The `SheetImport` for `path`: the Google Sheet when there is none,
    an `.xlsx` workbook, or a directory of CSVs."""
    if not path:
        return GoogleSheetImport(cfgs)
    if os.path.isdir(path):
        return CsvSheetImport(cfgs, path)
    if path.lower().endswith(".xlsx"):
        if not os.path.exists(path):
            raise ValueError(f"{path} does not exist")
        return XlsxSheetImport(cfgs, path)
    raise ValueError(f"{path}: give an .xlsx workbook or a directory of <tab name>.csv files")
